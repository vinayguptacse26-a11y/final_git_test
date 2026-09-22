from pathlib import Path
import logging
from pptx import Presentation

logger = logging.getLogger(__name__)


def load_ppt(file_path: Path) -> list[dict]:
 
    slides: list[dict] = []

    try:
        presentation = Presentation(file_path)

        for index, slide in enumerate(presentation.slides, start=1):
            texts = []
            has_image = False

            for shape in slide.shapes:
                
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())

               
                if shape.shape_type == 13:  # it means it has a picture
                    has_image = True

            slides.append({
                "slide_number": index,
                "text": "\n".join(texts),
                "has_image": has_image
            })

        logger.info(
            "Loaded PPT file: %s with %d slides",
            file_path.name,
            len(slides),
        )
        return slides

    except Exception as exc:
        logger.exception("Failed to load PPT file: %s", file_path)
        raise RuntimeError("PPT loading failed") from exc
